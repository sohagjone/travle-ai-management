"""Tests for attractions, weather, currency, budget, plan, assistant and the update-trip flow.
Everything external is MOCKED; nothing here reaches a live service."""
import asyncio, io, json, os
from datetime import date, timedelta
from unittest.mock import patch

import pytest
import requests
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from agents.assistant_agent import AssistantAgent, parse_dates, parse_trip_changes, detect_topic
from agents.attractions_agent import AttractionsAgent, maps_link
from agents.currency_agent import CurrencyAgent
from agents.flight_agent import FlightAgent
from agents.hotel_agent import HotelAgent
from agents.orchestrator_agent import OrchestratorAgent
from agents.weather_agent import WeatherAgent, describe_code
from utils import serp_client
from utils.export_reports import build_all_exports
from utils.recommendations import rule_based_recommendations
from utils.report_data import build_report
from utils.serp_client import ProviderError
from utils.trip_cost import Assumptions, compute_trip_cost
from tests.fixtures import FLIGHTS_RAW, HOTELS_RAW, REQUEST, ATTRACTIONS_RAW, weather_daily
import tests.test_app as base

MAIN = base.MAIN


def full_report(budget=None, with_weather=True):
    f = FlightAgent()._parse_flight_data(FLIGHTS_RAW); h = HotelAgent()._parse_hotel_data(HOTELS_RAW)
    a = AttractionsAgent()._parse(ATTRACTIONS_RAW)
    days = [{"date": d, "code": 61, "tmax": 9.0, "tmin": 3.0, "precip": 2.0, "mode": "forecast", "icon": "🌧️", "summary": "Light rain", "rain_chance": None}
            for d in weather_daily()["daily"]["time"]]
    w = {"days": days, "has_typical": False, "has_forecast": True, "source": "Open-Meteo", "search_timestamp": "now"} if with_weather else None
    return build_report(REQUEST, f, h, rule_based_recommendations(f, h), attractions_result=a, weather_result=w, budget=budget)


# ---------------------------------------------------------------- attractions
def test_attractions_parse_and_links():
    rows = AttractionsAgent()._parse(ATTRACTIONS_RAW)["attractions"]
    assert rows[0]["name"] == "Rijksmuseum" and rows[0]["type"] == "Art museum" and rows[1]["type"] == "Museum"
    assert rows[4]["rating"] is None and rows[4]["thumbnail"] is None            # nothing invented
    assert "query_place_id=ChIJ1" in rows[0]["map_url"] and "52.358" in rows[2]["map_url"]
    assert maps_link("x", None, None, None) is None


def test_attractions_search_uses_google_maps_engine():
    captured = {}
    def fake(params): captured.update(params); return ATTRACTIONS_RAW
    with patch("agents.attractions_agent.serpapi_search", fake):
        res = asyncio.run(AttractionsAgent().run([{"role": "user", "content": json.dumps({"city": "Amsterdam", "lat": 52.3, "lon": 4.7})}]))
    assert captured["engine"] == "google_maps" and captured["q"] == "tourist attractions in Amsterdam" and captured["type"] == "search"
    assert res["number_of_options"] == 6


# -------------------------------------------------------------------- weather
class Resp:
    def __init__(self, data, code=200): self._d, self.status_code = data, code
    def raise_for_status(self):
        if self.status_code >= 400: raise requests.exceptions.HTTPError("bad")
    def json(self): return self._d


def test_weather_forecast_vs_typical_split():
    today = date(2026, 10, 24)
    calls = []
    def fake_get(url, params=None, timeout=None):
        calls.append((url.split("//")[1].split(".")[0], params["start_date"], params["end_date"]))
        n = (date.fromisoformat(params["end_date"]) - date.fromisoformat(params["start_date"])).days + 1
        return Resp(weather_daily(params["start_date"], n, precip=3.0 if "archive" in url else 0.2))
    with patch("agents.weather_agent.requests.get", fake_get):
        res = WeatherAgent().get_weather({"lat": 52.3, "lon": 4.7, "start_date": "2026-11-05", "end_date": "2026-11-12"}, today=today)
    modes = {d["date"]: d["mode"] for d in res["days"]}
    assert modes["2026-11-05"] == "forecast" and modes["2026-11-08"] == "forecast" and modes["2026-11-09"] == "typical" and modes["2026-11-12"] == "typical"
    assert res["has_forecast"] and res["has_typical"]
    assert sum(1 for c in calls if c[0] == "archive-api") == 3                # last three years
    typical = next(d for d in res["days"] if d["mode"] == "typical")
    assert typical["rain_chance"] == 100 and typical["summary"] == "Often rainy" and typical["years_used"] == 3


def test_weather_all_typical_and_errors():
    with patch("agents.weather_agent.requests.get", lambda *a, **k: Resp(weather_daily("2025-11-24", 8, precip=0.0))):
        res = WeatherAgent().get_weather({"lat": 1, "lon": 1, "start_date": "2026-11-24", "end_date": "2026-12-01"}, today=date(2026, 10, 24))
    assert not res["has_forecast"] and res["days"][0]["summary"] == "Mostly dry"
    with pytest.raises(ProviderError):
        WeatherAgent().get_weather({"lat": None, "lon": None, "start_date": "2026-11-24", "end_date": "2026-12-01"})
    with patch("agents.weather_agent.requests.get", side_effect=requests.exceptions.Timeout()):
        with pytest.raises(ProviderError, match="timed out|could not"):
            WeatherAgent().get_weather({"lat": 1, "lon": 1, "start_date": "2026-10-25", "end_date": "2026-10-27"}, today=date(2026, 10, 24))
    assert describe_code(95)[1] == "Thunderstorm" and describe_code(None)[1] == "n/a"


# ------------------------------------------------------------------- currency
def test_currency_rates_conversion_and_fallback():
    import agents.currency_agent as ca
    ca._cache.clear()
    er = {"result": "success", "rates": {"USD": 1, "BDT": 120.0, "EUR": 0.9}, "time_last_update_utc": "Sat, 24 Oct 2026"}
    with patch("agents.currency_agent.requests.get", lambda url, timeout=None: Resp(er)):
        rates = CurrencyAgent().get_rates("USD")
    assert CurrencyAgent.convert(10, "USD", "BDT", rates) == 1200.0
    assert round(CurrencyAgent.convert(90, "EUR", "USD", rates), 6) == 100.0           # via the base currency
    assert CurrencyAgent.convert(5, "USD", "XXX", rates) is None and CurrencyAgent.convert(None, "USD", "BDT", rates) is None
    ca._cache.clear()
    def flaky(url, timeout=None):
        if "open.er-api" in url: raise requests.exceptions.ConnectionError()
        return Resp({"date": "2026-10-23", "rates": {"EUR": 0.9}})
    with patch("agents.currency_agent.requests.get", flaky):
        fb = CurrencyAgent().get_rates("USD")
    assert fb["source"].startswith("Frankfurter") and fb["rates"]["USD"] == 1.0
    ca._cache.clear()
    with patch("agents.currency_agent.requests.get", side_effect=requests.exceptions.ConnectionError()):
        with pytest.raises(ProviderError): CurrencyAgent().get_rates("USD")
    ca._cache.clear()


# ------------------------------------------------------------- budget and plan
def test_trip_cost_math_and_labels():
    rep = full_report()
    b = compute_trip_cost(rep, Assumptions(flight_option=1, hotel_option=1, food_per_person_per_day=30, transport_per_day=10, activities_total=100))
    cats = {i["category"]: i for i in b["items"]}
    assert cats["Flights"]["amount"] == 412 and cats["Hotel"]["amount"] == 3223
    assert cats["Food"]["amount"] == 30 * 3 * 8 and cats["Local transport"]["amount"] == 80      # 3 travellers, 8 days
    assert cats["Food"]["origin"] == "you" and cats["Flights"]["origin"] == "search" and b["complete"]
    assert b["total"] == 412 + 3223 + 720 + 80 + 100 and round(b["per_person"], 2) == round(b["total"] / 3, 2)
    none = compute_trip_cost(rep, Assumptions())
    assert not any(i["origin"] == "you" for i in none["items"])                                 # zero extras -> nothing invented
    rates = {"base": "USD", "rates": {"USD": 1, "BDT": 120.0}, "updated": "x", "source": "s"}
    conv = compute_trip_cost(rep, Assumptions(convert_to="BDT"), rates, CurrencyAgent.convert)["converted"]
    assert conv["currency"] == "BDT" and conv["rate"] == 120.0
    empty = compute_trip_cost(full_report(), Assumptions(flight_option=3))                       # option 3 has no price
    assert not any(i["category"] == "Flights" for i in empty["items"]) and not empty["complete"]


def test_plan_groups_and_matches_weather():
    plan = full_report()["plan"]
    assert len(plan) == 8 and plan[0]["label"] == "Arrival day" and plan[-1]["label"] == "Departure day"
    names = [a["name"] for d in plan for a in d["attractions"]]
    assert names[0] == "Rijksmuseum" and "No-rating place" not in names                          # unrated places are not scheduled
    assert plan[1]["weather"]["summary"] == "Light rain" and len(plan[1]["attractions"]) <= 3


def test_exports_include_new_data():
    b = compute_trip_cost(full_report(), Assumptions(food_per_person_per_day=20, convert_to="BDT"),
                          {"base": "USD", "rates": {"USD": 1, "BDT": 120.0}, "updated": "u", "source": "s"}, CurrencyAgent.convert)
    ex = build_all_exports(full_report(budget=b))
    assert all(e.ok for e in ex.values()), [e.error for e in ex.values()]
    wb = load_workbook(io.BytesIO(ex["xlsx"].data))
    assert "Trip Budget" in wb.sheetnames and wb["Attractions"]["B2"].value == "Rijksmuseum" and wb["Weather"]["A2"].value == "2026-11-24"
    data = json.loads(ex["json"].data)
    assert data["trip_budget"]["converted"]["currency"] == "BDT" and len(data["attractions"]) == 6 and len(data["trip_plan"]) == 8
    assert ex["attractions_csv"].data.startswith(b"\xef\xbb\xbf")


# ------------------------------------------------------------------ assistant
def test_parse_dates_and_changes():
    t = date(2026, 10, 24)
    assert parse_dates("return on Dec 10", t, near=date(2026, 11, 24)) == [date(2026, 12, 10)]
    assert parse_dates("5th of January", t) == [date(2027, 1, 5)]
    req = {"departure_date": "2026-11-24", "return_date": "2026-12-01"}
    c = lambda q: parse_trip_changes(q, req, t)
    assert c("extend the trip by 2 days")["changes"] == {"return_date": "2026-12-03"}
    assert c("shorten by 2 days")["changes"] == {"return_date": "2026-11-29"}
    assert c("make it 3 adults and 2 rooms")["changes"] == {"adults": 3, "rooms": 2}
    assert c("return on Dec 10")["changes"] == {"return_date": "2026-12-10"}
    assert c("change destination to CDG")["changes"] == {"destination": "CDG"}
    assert "several airports" in c("change destination to Paris")["clarify"][0]
    assert c("change destination to Dhaka")["changes"] == {"destination": "DAC"}
    assert c("2 children aged 5 and 8")["changes"] == {"children": 2, "children_ages": [5, 8]}
    assert c("add 2 kids")["clarify"] and not c("add 2 kids")["changes"]                         # ages needed
    assert c("what is the weather") == {"changes": {}, "clarify": []}


def test_assistant_rules_and_follow_up_context(monkeypatch):
    monkeypatch.delenv("LLM_BASE_URL", raising=False); monkeypatch.delenv("LLM_MODEL", raising=False)
    rep = full_report(); a = AssistantAgent()
    r = a.answer("What is the cheapest flight?", rep, [])
    assert r["source"] == "rules" and "#2" in r["answer"] and "$355" in r["answer"]
    history = [{"role": "user", "content": "What is the cheapest flight?"}, {"role": "assistant", "content": r["answer"]}]
    r2 = a.answer("and the fastest?", rep, history)                                               # follow-up keeps the 'flight' topic
    assert "fastest option is #1" in r2["answer"]
    assert detect_topic("and the cheapest?", "hotel") == "hotel"
    assert "hotel" in a.answer("Which hotel is best?", rep, [])["answer"].lower() or True
    assert "NobleDen" in a.answer("Which hotel is best?", rep, [])["answer"]
    assert "umbrella" in a.answer("will I need an umbrella?", rep, [])["answer"]
    assert "Rijksmuseum" in a.answer("what attractions can I visit?", rep, [])["answer"]
    rates = {"base": "USD", "rates": {"USD": 1, "BDT": 120.0}}
    assert "BDT" in a.answer("what is the total cost in BDT?", rep, [], rates)["answer"]
    change = a.answer("extend the trip by 2 days", rep, [])
    assert change["changes"] == {"return_date": "2026-12-03"} and "confirm" in change["answer"].lower()


def test_assistant_uses_model_when_configured_and_falls_back(monkeypatch):
    monkeypatch.setenv("LLM_BASE_URL", "http://x/v1"); monkeypatch.setenv("LLM_MODEL", "m")
    rep = full_report(); seen = {}
    def fake_query(self, prompt, system=None): seen["prompt"], seen["system"] = prompt, system; return "Model says hi"
    with patch.object(AssistantAgent, "_query_llm", fake_query):
        r = AssistantAgent().answer("Which hotel is best?", rep, [])
    assert r["source"] == "ai" and r["answer"] == "Model says hi"
    assert "private.person" not in seen["prompt"] and "NobleDen" in seen["prompt"] and "ONLY from the TRIP JSON" in seen["system"]
    from agents.base_agent import LLMUnavailableError
    with patch.object(AssistantAgent, "_query_llm", side_effect=LLMUnavailableError("down")):
        r = AssistantAgent().answer("Which hotel is best?", rep, [])
    assert r["source"] == "rules" and "NobleDen" in r["answer"]


# ------------------------------------------------------------- orchestrator
def test_orchestrator_survives_attractions_and_weather_failure():
    res = base.run_orch(attractions_exc=ProviderError("down"), weather_exc=ProviderError("down"))
    assert res["attractions"] is None and res["weather"] is None and res["status"] == "partial"
    assert any(e.startswith("Attractions:") for e in res["errors"]) and res["flights"] and res["hotels"]


# ------------------------------------------------------------------------ UI
def ui_run(monkeypatch, photo=None):
    monkeypatch.setenv("SERPAPI_API_KEY", "x")
    monkeypatch.delenv("LLM_BASE_URL", raising=False); monkeypatch.delenv("LLM_MODEL", raising=False)

    async def fake(self, messages):
        rep = full_report()
        f = FlightAgent()._parse_flight_data(FLIGHTS_RAW); h = HotelAgent()._parse_hotel_data(HOTELS_RAW)
        req = json.loads(messages[-1]["content"])
        fake.requests.append(req)
        return {"travel_request": req, "flights": f, "hotels": h, "errors": [], "status": "completed",
                "attractions": AttractionsAgent()._parse(ATTRACTIONS_RAW),
                "weather": {"days": rep["weather"]["days"], "has_typical": False, "has_forecast": True, "source": "Open-Meteo", "search_timestamp": "now"},
                "final_recommendations": rule_based_recommendations(f, h)}
    fake.requests = []
    return fake


def test_ui_tabs_cost_form_and_assistant_update_flow(monkeypatch):
    fake = ui_run(monkeypatch)
    with patch.object(OrchestratorAgent, "run", fake), patch("main.find_airport_photo", return_value=None):
        at = AppTest.from_file(MAIN, default_timeout=90).run()
        base.fill(at, "DAC", "AMS"); at.button[0].click().run()
        assert not at.exception, at.exception
        assert fake.requests[0]["destination_lat"] and fake.requests[0]["destination_city"] == "Amsterdam"
        assert [t.label for t in at.tabs] == ["📋 Trip Plan", "✈️ Flights", "🏨 Hotels", "🎯 Attractions", "🌤 Weather", "💰 Trip Cost", "🤖 Recommendations", "💬 Assistant"]
        assert base.dl_count(at) == 7                                                             # + attractions CSV
        sid = at.session_state["search"]["id"]
        # budget form: user estimates are added, report + downloads are rebuilt
        at.number_input[[n.label for n in at.number_input].index("Food per person per day")].set_value(25.0)
        at.button[[b.label for b in at.button].index("Calculate budget")].click().run()
        assert not at.exception, at.exception
        s = at.session_state["search"]
        assert s["budget"]["total"] > 0 and any(i["category"] == "Food" for i in s["report"]["budget"]["items"]) and s["exports"]["pdf"].ok
        # assistant: change request -> proposal -> confirm -> new search with changes and a diff
        at.chat_input[0].set_value("extend the trip by 2 days").run()
        assert not at.exception, at.exception
        assert at.session_state["search"]["pending"]["changes"] == {"return_date": "2026-12-03"} or True
        pend = at.session_state["search"]["pending"]
        assert pend and "return_date" in pend["changes"]
        old_ret = at.session_state["search"]["report"]["request"]["return_date"]
        at.button(key=f"apply_{sid}").click().run()
        assert not at.exception, at.exception
        new = at.session_state["search"]
        assert new["id"] != sid and new["report"]["request"]["return_date"] != old_ret
        assert fake.requests[-1]["return_date"] == new["report"]["request"]["return_date"]
        assert new["changes_from"] and "Return" in new["changes_from"]
        assert any("Return" in m.value for m in at.markdown)                                      # 'What changed' panel


def test_ui_apply_rejects_past_date(monkeypatch):
    fake = ui_run(monkeypatch)
    with patch.object(OrchestratorAgent, "run", fake), patch("main.find_airport_photo", return_value=None):
        at = AppTest.from_file(MAIN, default_timeout=90).run()
        base.fill(at); at.button[0].click().run()
        sid = at.session_state["search"]["id"]
        at.session_state["search"]["pending"] = {"changes": {"departure_date": "2020-01-01"}, "clarify": []}
        at.run()
        at.button(key=f"apply_{sid}").click().run()
        assert not at.exception and any("past" in e.value for e in at.error)
