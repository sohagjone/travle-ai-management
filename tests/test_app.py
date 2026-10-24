"""Tests. ALL provider calls are mocked (fixtures); nothing here touches live APIs."""
import asyncio, csv, io, json, zipfile
from datetime import date, timedelta
from unittest.mock import patch

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from docx import Document
from streamlit.testing.v1 import AppTest

from agents.flight_agent import FlightAgent
from agents.hotel_agent import HotelAgent
from agents.orchestrator_agent import OrchestratorAgent
from agents.email_agent import EmailAgent
from agents.attractions_agent import AttractionsAgent
from agents.weather_agent import WeatherAgent
from utils import serp_client
from utils.export_reports import build_all_exports, csv_safe
from utils.recommendations import rule_based_recommendations, parse_llm_recommendations
from utils.report_data import build_report
from utils.validation import validate_travel_request
import os
MAIN = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'main.py')
from tests.fixtures import FLIGHTS_RAW, HOTELS_RAW, REQUEST, ATTRACTIONS_RAW, weather_daily


def make_report(flights=FLIGHTS_RAW, hotels=HOTELS_RAW, request=REQUEST):
    f = FlightAgent()._parse_flight_data(flights) if flights is not None else None
    h = HotelAgent()._parse_hotel_data(hotels) if hotels is not None else None
    return build_report(request, f, h, rule_based_recommendations(f, h))


# ---------- parsing / validation
def test_flight_parsing_one_row_per_itinerary_and_best_flights_included():
    rows = FlightAgent()._parse_flight_data(FLIGHTS_RAW)["recommended_flights"]
    assert len(rows) == 3 and rows[0]["is_top_pick"] and rows[1]["airline"] == "Air France, KLM"
    assert rows[2]["price_value"] is None and rows[2]["duration"] is None  # nothing invented


def test_hotel_parsing_handles_missing_fields():
    rows = HotelAgent()._parse_hotel_data(HOTELS_RAW)["recommended_hotels"]
    assert rows[1]["rating"] is None and rows[1]["amenities"] == [] and rows[2]["image_url"] is None


def test_validation():
    d = date.today() + timedelta(days=10)
    assert validate_travel_request("MAD", "AMS", d, d + timedelta(days=3), 1, 0, 1) == []
    assert any("later" in e for e in validate_travel_request("MAD", "AMS", d, d, 1, 0, 1))
    assert any("past" in e for e in validate_travel_request("MAD", "AMS", date(2000, 1, 1), d, 1, 0, 1))
    assert len(validate_travel_request("", "", d, d + timedelta(days=1), 0, 0, 0)) >= 3
    assert any("age" in e for e in validate_travel_request("MAD", "AMS", d, d + timedelta(days=2), 2, 1, 1, []))


# ---------- exports
def test_all_exports_valid_and_no_email_leak():
    ex = build_all_exports(make_report())
    assert all(e.ok for e in ex.values()), [e.error for e in ex.values()]
    pdf = PdfReader(io.BytesIO(ex["pdf"].data)); text = "".join(p.extract_text() for p in pdf.pages)
    assert "NobleDen" in text and "Page 1 of" in text
    wb = load_workbook(io.BytesIO(ex["xlsx"].data))
    assert wb.sheetnames == ["Search Summary", "Flight Options", "Hotel Options", "Attractions", "Weather", "Trip Plan", "Recommendations", "Sources and Notes"]
    fl = wb["Flight Options"]
    assert fl["J2"].value == 412 and fl["L2"].hyperlink is not None
    assert fl["B4"].data_type == "s" and fl["B4"].value.startswith("=")  # stays text, not a formula
    assert Document(io.BytesIO(ex["docx"].data)).paragraphs
    rows = list(csv.reader(io.StringIO(ex["hotels_csv"].data.decode("utf-8-sig"))))
    assert rows[0][1] == "Hotel" and "東京" in ex["hotels_csv"].data.decode("utf-8-sig")
    assert ex["hotels_csv"].data.startswith(b"\xef\xbb\xbf")
    data = json.loads(ex["json"].data)
    assert data["search_request"]["adults"] == 2 and len(data["flights"]) == 3
    for e in ex.values():
        blob = e.data
        if e.key in ("xlsx", "docx"):
            blob = b"".join(zipfile.ZipFile(io.BytesIO(e.data)).read(n) for n in zipfile.ZipFile(io.BytesIO(e.data)).namelist())
        if e.key != "pdf":
            assert b"private.person" not in blob


def test_empty_results_exports_still_valid():
    for f, h in ((None, HOTELS_RAW), (FLIGHTS_RAW, None), ({}, {})):
        ex = build_all_exports(make_report(f if f is not None else {}, h if h is not None else {}))
        assert all(e.ok for e in ex.values())
    ex = build_all_exports(make_report({}, {}))
    assert json.loads(ex["json"].data)["cost_estimate"]["scenarios"] == []


def test_csv_formula_injection_guard():
    assert csv_safe("=1+1") == "'=1+1" and csv_safe("ok") == "ok"


def test_filename():
    assert build_all_exports(make_report())["pdf"].filename.startswith("travel_report_MAD_to_AMS_")


# ---------- recommendations / orchestrator
def test_llm_output_validation():
    assert parse_llm_recommendations({"error": "x"}, 3, 3) is None
    assert parse_llm_recommendations({"best_flight_option": 99, "best_flight_reason": "x"}, 3, 3) is None
    ok = parse_llm_recommendations({"best_flight_option": 2, "best_flight_reason": "r", "best_hotel_option": None}, 3, 3)
    assert ok["best_flight"]["option"] == 2 and ok["best_hotel"] is None


def run_orch(flight_exc=None, hotel_exc=None, env=None, attractions_exc=None, weather_exc=None):
    async def ff(self, m): 
        if flight_exc: raise flight_exc
        return FlightAgent()._parse_flight_data(FLIGHTS_RAW)
    async def hh(self, m):
        if hotel_exc: raise hotel_exc
        return HotelAgent()._parse_hotel_data(HOTELS_RAW)
    async def aa(self, m):
        if attractions_exc: raise attractions_exc
        return AttractionsAgent()._parse(ATTRACTIONS_RAW)
    async def ww(self, m):
        if weather_exc: raise weather_exc
        return {"days": [{**r, "mode": "forecast", "icon": "x", "summary": "Rain", "tmax": r["tmax"], "tmin": r["tmin"], "precip": r["precip"], "rain_chance": None} for r in [
            {"date": d, "code": 61, "tmax": 9.0, "tmin": 3.0, "precip": 2.0} for d in weather_daily()["daily"]["time"]]],
                "unit": "°C", "has_typical": False, "has_forecast": True, "source": "Open-Meteo", "search_timestamp": "now"}
    with patch.object(FlightAgent, "run", ff), patch.object(HotelAgent, "run", hh), \
         patch.object(AttractionsAgent, "run", aa), patch.object(WeatherAgent, "run", ww), \
         patch.dict("os.environ", env or {}, clear=False):
        return asyncio.run(OrchestratorAgent().run([{"role": "user", "content": json.dumps(REQUEST)}]))


def test_orchestrator_no_llm_falls_back_and_never_sends_email():
    with patch.object(EmailAgent, "send_email", side_effect=AssertionError("email must not be sent")):
        res = run_orch(env={"LLM_BASE_URL": "", "LLM_MODEL": ""})
    assert res["final_recommendations"]["source"] == "rule_based" and res["status"] == "completed"


def test_orchestrator_partial_and_total_failure():
    res = run_orch(flight_exc=serp_client.ProviderError("boom"))
    assert res["status"] == "partial" and res["flights"] is None and res["hotels"]
    with pytest.raises(RuntimeError):
        run_orch(serp_client.ProviderError("a"), serp_client.ProviderError("b"))


def test_unreachable_llm_does_not_crash():
    res = run_orch(env={"LLM_BASE_URL": "http://127.0.0.1:9/v1", "LLM_MODEL": "x", "LLM_API_KEY": "k"})
    assert res["final_recommendations"]["source"] == "rule_based"
    assert any("unavailable" in w for w in res["final_recommendations"]["warnings"])


# ---------- provider client: missing key / errors never leak the key
def test_missing_key(monkeypatch):
    monkeypatch.delenv("SERPAPI_API_KEY", raising=False)
    with pytest.raises(serp_client.ProviderError, match="not configured"):
        serp_client.serpapi_search({"engine": "google_flights"})


class FakeResp:
    def __init__(self, code, payload): self.status_code, self._p = code, payload
    def json(self):
        if self._p is None: raise ValueError
        return self._p


@pytest.mark.parametrize("code,payload,expect", [
    (401, {"error": "bad"}, "rejected"), (429, {"error": "quota"}, "rate limit"),
    (500, None, "server error"), (200, {"error": "Google Flights hasn't returned any results"}, None)])
def test_provider_errors(monkeypatch, code, payload, expect):
    monkeypatch.setenv("SERPAPI_API_KEY", "SECRETKEY123")
    with patch("utils.serp_client.requests.get", return_value=FakeResp(code, payload)):
        if expect is None:
            with pytest.raises(serp_client.NoResultsError): serp_client.serpapi_search({"engine": "x"})
        else:
            with pytest.raises(serp_client.ProviderError, match=expect) as e:
                serp_client.serpapi_search({"engine": "x"})
            assert "SECRETKEY123" not in str(e.value)


def test_timeout_message(monkeypatch):
    import requests
    monkeypatch.setenv("SERPAPI_API_KEY", "SECRETKEY123")
    with patch("utils.serp_client.requests.get", side_effect=requests.exceptions.Timeout()):
        with pytest.raises(serp_client.ProviderError, match="timed out"): serp_client.serpapi_search({"engine": "x"})


# ---------- email
def test_email_escapes_html_and_requires_real_success():
    report = make_report()
    html = EmailAgent()._generate_email_content({"report": report})
    assert "<b>Inn</b>" not in html and "&lt;b&gt;Inn" in html and "javascript:" not in html
    r = asyncio.run(EmailAgent().run([{"role": "user", "content": json.dumps({"to_email": "bad", "report": report})}]))
    assert r["success"] is False
    with patch.dict("os.environ", {"SENDGRID_KEY": "", "SENDGRID_FROM_EMAIL": ""}):
        ok, err = EmailAgent().send_email("a@b.com", "x")
    assert ok is False and "not configured" in err


# ---------- Streamlit UI (AppTest with mocked search)
def fake_run(self, messages):
    async def go(): return {"travel_request": {}, "flights": FlightAgent()._parse_flight_data(FLIGHTS_RAW),
                            "hotels": HotelAgent()._parse_hotel_data(HOTELS_RAW), "errors": [], "status": "completed",
                            "final_recommendations": rule_based_recommendations(FlightAgent()._parse_flight_data(FLIGHTS_RAW), HotelAgent()._parse_hotel_data(HOTELS_RAW))}
    return go()


from utils.validation import airport_label


def fill(at, o="MAD", d="AMS"):
    at.selectbox(key="origin").select(airport_label(o)); at.selectbox(key="destination").select(airport_label(d))
    return at


def dl_count(at): return len(at.get("download_button"))


def test_ui_flow(monkeypatch):
    monkeypatch.setenv("SERPAPI_API_KEY", "x")
    with patch.object(OrchestratorAgent, "run", fake_run):
        at = AppTest.from_file(MAIN, default_timeout=60).run()
        assert not at.exception and dl_count(at) == 0           # no buttons before a search
        fill(at, "MAD", "MAD"); at.button[0].click().run()      # invalid: same airports
        assert at.error and dl_count(at) == 0
        fill(at); at.button[0].click().run()
        assert not at.exception and dl_count(at) == 6            # all six downloads appear
        at.button(key="send_" + at.session_state["search"]["id"]) if False else None
        at.run()                                                  # plain rerun keeps results + buttons
        assert dl_count(at) == 6
        first_id = at.session_state["search"]["id"]
        fill(at, "MAD", "BCN"); at.button[0].click().run()        # new search replaces old
        assert at.session_state["search"]["id"] != first_id
        assert at.session_state["search"]["report"]["request"]["destination"] == "BCN"
        [b for b in at.button if b.label == "New search"][0].click().run()
        assert "search" not in at.session_state and dl_count(at) == 0


def test_ui_total_failure_shows_error_no_buttons(monkeypatch):
    monkeypatch.setenv("SERPAPI_API_KEY", "x")
    async def boom(self, m): raise RuntimeError("Flights: SERPAPI_API_KEY is not configured.")
    with patch.object(OrchestratorAgent, "run", boom):
        at = fill(AppTest.from_file(MAIN, default_timeout=60).run()); at.button[0].click().run()
        assert at.error and dl_count(at) == 0


def test_airport_dropdown_labels():
    from utils.validation import airport_choices, code_from_label
    labels, mapping = airport_choices()
    assert "DAC-Hazrat Shahjalal International Airport, Bangladesh" in labels and len(labels) > 5000
    assert code_from_label("DAC-Hazrat Shahjalal International Airport, Bangladesh") == "DAC" and code_from_label(None) == ""


def test_ui_requires_airports_and_shows_cards(monkeypatch):
    monkeypatch.setenv("SERPAPI_API_KEY", "x")
    with patch.object(OrchestratorAgent, "run", fake_run):
        at = AppTest.from_file(MAIN, default_timeout=60).run()
        assert at.selectbox(key="origin").value is None
        at.button[0].click().run()
        assert at.error and dl_count(at) == 0                       # nothing selected -> friendly errors
        fill(at); at.button[0].click().run()
        assert not at.exception
        assert {"Flights found", "Cheapest", "Fastest"} <= {m.label for m in at.metric}
        text = " ".join(m.value for m in at.markdown)
        assert "Cheapest" in text and "View on Google Flights" in text
        at.selectbox(key="sort_flights").select("Duration (shortest)").run()
        assert not at.exception


# ---------- filters
def _rows():
    f = FlightAgent()._parse_flight_data(FLIGHTS_RAW)["recommended_flights"]
    h = make_report()["hotels"]
    return make_report()["flights"], h


def test_filter_flights():
    from utils.display import filter_flights, airline_options
    flights, _ = _rows()
    assert airline_options(flights) == sorted({"Iberia", "Air France", "KLM", '=HYPERLINK("http://evil","x")'})
    assert len(filter_flights(flights)) == 3                                   # no filter keeps everything (even unpriced)
    assert [f["option"] for f in filter_flights(flights, (300, 400), (355, 412))] == [2]
    assert [f["option"] for f in filter_flights(flights, airlines=["KLM"])] == [2]  # connecting itinerary includes KLM
    assert [f["option"] for f in filter_flights(flights, max_stops=0)] == [1, 3]
    assert [f["option"] for f in filter_flights(flights, depart_hours=(0, 7))] == [2]


def test_filter_hotels():
    from utils.display import filter_hotels, amenity_options
    _, hotels = _rows()
    assert len(filter_hotels(hotels)) == 3
    assert [h["option"] for h in filter_hotels(hotels, min_rating=4.5)] == [1]
    assert [h["option"] for h in filter_hotels(hotels, (100, 200), (150, 537))] == [3]
    assert [h["option"] for h in filter_hotels(hotels, amenities=["Parking"])] == [1]
    assert [h["option"] for h in filter_hotels(hotels, name_query="canal")] == [3]
    assert "Free Wi-Fi" in amenity_options(hotels)


# ---------- airport photos (Wikipedia is mocked; the sandbox cannot reach it)
class WikiResp:
    def __init__(self, pages): self._pages = pages
    def raise_for_status(self): pass
    def json(self): return {"query": {"pages": self._pages}}


def test_airport_photo_matching_and_safety():
    import requests
    from utils import airport_photos as ap
    good = {"1": {"title": "Hazrat Shahjalal International Airport", "index": 2, "thumbnail": {"source": "https://upload.wikimedia.org/x.jpg"}},
            "2": {"title": "Dhaka", "index": 1, "thumbnail": {"source": "https://upload.wikimedia.org/city.jpg"}}}
    with patch("utils.airport_photos.requests.get", return_value=WikiResp(good)):
        p = ap.find_airport_photo("DAC")
        assert p["url"] == "https://upload.wikimedia.org/x.jpg" and p["credit"] == "Wikipedia"   # skips unrelated "Dhaka"
    evil = {"1": {"title": "Hazrat Shahjalal International Airport", "index": 1, "thumbnail": {"source": "http://evil.example/x.jpg"}}}
    with patch("utils.airport_photos.requests.get", return_value=WikiResp(evil)), \
         patch("utils.airport_photos.serpapi_search", side_effect=serp_client.ProviderError("no key")):
        assert ap.find_airport_photo("DAC") is None                                  # non-Wikimedia / non-https blocked
    assert ap.find_airport_photo("ZZZ") is None


def test_airport_photo_fallback_chain():
    import requests
    from utils import airport_photos as ap
    # Wikipedia blocked (e.g. 403/429 from a shared cloud IP) -> Google Images via SerpAPI is used
    serp = {"images_results": [
        {"title": "Unrelated beach", "thumbnail": "https://encrypted-tbn0.gstatic.com/a.jpg", "link": "https://x.test/1"},
        {"title": "Hazrat Shahjalal International Airport terminal", "thumbnail": "https://encrypted-tbn0.gstatic.com/b.jpg", "link": "https://x.test/2"},
        {"title": "Hazrat Shahjalal evil", "thumbnail": "http://evil.example/c.jpg", "link": "https://x.test/3"}]}
    with patch("utils.airport_photos.requests.get", side_effect=requests.exceptions.Timeout()), \
         patch("utils.airport_photos.serpapi_search", return_value=serp):
        p = ap.find_airport_photo("DAC")
    assert p["url"].endswith("/b.jpg") and p["credit"] == "Google Images"
    # everything fails -> None (placeholder), never an exception
    with patch("utils.airport_photos.requests.get", side_effect=requests.exceptions.Timeout()), \
         patch("utils.airport_photos.serpapi_search", side_effect=serp_client.ProviderError("quota")):
        assert ap.find_airport_photo("DAC") is None


def test_ui_tiles_filters_and_no_more_options(monkeypatch):
    monkeypatch.setenv("SERPAPI_API_KEY", "x")
    photo = {"url": "https://upload.wikimedia.org/x.jpg", "title": "t", "page": "https://en.wikipedia.org/wiki/t"}
    with patch.object(OrchestratorAgent, "run", fake_run), patch("main.find_airport_photo", return_value=photo):
        at = AppTest.from_file(MAIN, default_timeout=60).run()
        assert not at.exception and not at.expander or all("More options" not in e.label for e in at.expander)
        assert not [s for s in at.selectbox if s.label == "Flight stops"]
        fill(at); at.button[0].click().run()
        assert not at.exception
        assert at.session_state["search"]["report"]["request"]["flight_stops_filter"] == "Any number of stops"
        assert len(at.metric) >= 5 or len([m for m in at.metric]) >= 3
        sid = at.session_state["search"]["id"]
        at.multiselect(key=f"flt_airlines_{sid}").select("KLM").run()
        assert not at.exception and any("Showing 1 of 3 flights" in c.value for c in at.caption)
        at.selectbox(key=f"flt_stops_{sid}").select("Nonstop only").run()
        assert any("Showing 0 of 3" in c.value for c in at.caption)      # KLM option is a 1-stop itinerary
        at.select_slider(key=f"hflt_rating_{sid}").set_value(4.5).run()
        assert any("Showing 1 of 3 hotels" in c.value for c in at.caption)
        assert dl_count(at) == 6                                          # downloads unaffected by filters


def test_airport_labels_have_country_and_are_unique():
    from utils.validation import airport_choices, airport_label, country_name
    labels, mapping = airport_choices()
    assert len(labels) == len(set(labels)) == len(mapping)
    assert airport_label("AMS").endswith(", Netherlands") and airport_label("JFK").endswith(", United States")
    assert airport_label("ICN").endswith(", South Korea") and country_name("ZZ") == "ZZ"
    assert all(", " in l for l in labels)
