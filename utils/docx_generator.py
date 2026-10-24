"""Editable Word (.docx) report built with python-docx. Same content as the PDF."""

from __future__ import annotations

import io
from typing import Any, Dict

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .report_data import (
    ATTRACTION_COLUMNS,
    FLIGHT_COLUMNS,
    HOTEL_COLUMNS,
    clean_text,
    display_value,
    fmt_money,
    safe_http_url,
)


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_page_number_footer(doc: Document) -> None:
    paragraph = doc.sections[0].footer.paragraphs[0]
    paragraph.text = "Search results only - not a booking. Verify prices before booking.   Page "
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    for r in paragraph.runs:
        r.font.size = Pt(8)


def _cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(clean_text(text))
    run.font.size = Pt(8)
    run.bold = bold


def build_docx(report: Dict[str, Any]) -> bytes:
    req = report["request"]
    currency = req.get("currency") or "USD"
    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Cm(1.8)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    doc.add_heading(clean_text(report["title"]), 0)
    doc.add_paragraph(f"{req.get('origin')} to {req.get('destination')} | {req.get('departure_date')} - {req.get('return_date')}")
    doc.add_paragraph(f"Generated {report['generated_at_display']}").runs[0].font.size = Pt(8)

    doc.add_heading("Search criteria", 1)
    travellers = f"{req.get('adults')} adult(s), {req.get('children')} child(ren)"
    if req.get("children_ages"):
        travellers += f" (ages {', '.join(str(a) for a in req['children_ages'])})"
    pairs = [
        ("Origin", req.get("origin")), ("Destination", req.get("destination")),
        ("Hotel search city", req.get("hotel_search_city") or "n/a"),
        ("Departure date", req.get("departure_date")), ("Return date", req.get("return_date")),
        ("Nights", req.get("nights") if req.get("nights") is not None else "n/a"),
        ("Travellers", travellers), ("Rooms", req.get("rooms")),
        ("Flight stops filter", req.get("flight_stops_filter")), ("Currency", currency),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for k, v in pairs:
        cells = table.add_row().cells
        _cell_text(cells[0], str(k), bold=True)
        _cell_text(cells[1], str(v))

    recs = report["recommendations"]
    doc.add_heading("Recommendations", 1)
    doc.add_paragraph(clean_text(recs.get("source_label") or "")).runs[0].italic = True
    shown = False
    for label, key, rows in (("Best flight", "best_flight", report["flights"]), ("Best hotel", "best_hotel", report["hotels"])):
        pick = recs.get(key)
        if pick:
            item = next((r for r in rows if r["option"] == pick["option"]), None)
            name = (item or {}).get("airline") or (item or {}).get("name") or ""
            p = doc.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(clean_text(f"option {pick['option']} - {name}. {pick.get('reason') or ''}"))
            shown = True
    if recs.get("summary"):
        doc.add_paragraph(clean_text(recs["summary"]))
        shown = True
    if not shown:
        doc.add_paragraph("No recommendations could be produced from the returned data.")
    for w in recs.get("warnings") or []:
        doc.add_paragraph(clean_text("Note: " + w)).runs[0].font.size = Pt(8)

    doc.add_heading("Estimated trip cost", 1)
    if not report["cost_estimate"]["scenarios"]:
        doc.add_paragraph("No estimate possible: no priced flight or hotel was returned.")
    for sc in report["cost_estimate"]["scenarios"]:
        doc.add_paragraph().add_run(sc["scenario"]).bold = True
        t = doc.add_table(rows=1, cols=3)
        t.style = "Light Grid Accent 1"
        for c, h in zip(t.rows[0].cells, ["Item", "Basis", f"Amount ({currency})"]):
            _cell_text(c, h, bold=True)
        for item in sc["items"]:
            cells = t.add_row().cells
            _cell_text(cells[0], item["label"]); _cell_text(cells[1], item["basis"]); _cell_text(cells[2], fmt_money(item["amount"], currency))
        cells = t.add_row().cells
        _cell_text(cells[0], "Total of items above" + ("" if sc["complete"] else " (INCOMPLETE: flight or hotel price missing)"), bold=True)
        _cell_text(cells[2], fmt_money(sc["total"], currency), bold=True)
    doc.add_paragraph("Not included: " + "; ".join(report["cost_estimate"]["excluded"]) + ".").runs[0].font.size = Pt(8)

    budget = report.get("budget")
    if budget and budget.get("items"):
        doc.add_heading("Trip budget", 1)
        t = doc.add_table(rows=1, cols=4)
        t.style = "Light Grid Accent 1"
        for c, h in zip(t.rows[0].cells, ["Category", "Item", "Source", f"Amount ({currency})"]):
            _cell_text(c, h, bold=True)
        for item in budget["items"]:
            cells = t.add_row().cells
            for cell, text in zip(cells, [item["category"], item["label"], item["basis"], fmt_money(item["amount"], currency)]):
                _cell_text(cell, text)
        cells = t.add_row().cells
        _cell_text(cells[0], "Total" + ("" if budget.get("complete") else " (INCOMPLETE)"), bold=True)
        _cell_text(cells[3], fmt_money(budget["total"], currency), bold=True)
        conv = budget.get("converted")
        if conv:
            doc.add_paragraph(f"Converted total: {fmt_money(conv['total'], conv['currency'])} (1 {currency} = {conv['rate']:.4f} {conv['currency']}; indicative rate, updated {conv.get('updated') or 'n/a'}).")

    doc.add_heading(f"Flight options ({len(report['flights'])})", 1)
    if not report["flights"]:
        doc.add_paragraph("No flight options were returned for this search.")
    else:
        keys = ["option", "airline", "flight_number", "departure", "arrival", "duration", "stops", "price"]
        cols = [c for c in FLIGHT_COLUMNS if c.key in keys]
        t = doc.add_table(rows=1, cols=len(cols) + 1)
        t.style = "Light Grid Accent 1"
        for c, h in zip(t.rows[0].cells, [c.header for c in cols] + ["Link"]):
            _cell_text(c, h, bold=True)
        for row in report["flights"]:
            cells = t.add_row().cells
            for cell, col in zip(cells, cols):
                _cell_text(cell, display_value(row, col))
            url = safe_http_url(row.get("booking_url"))
            cells[-1].text = ""
            if url:
                _add_hyperlink(cells[-1].paragraphs[0], url, "Google Flights")
            else:
                _cell_text(cells[-1], "n/a")

    doc.add_heading(f"Hotel options ({len(report['hotels'])})", 1)
    if not report["hotels"]:
        doc.add_paragraph("No hotel options were returned for this search.")
    else:
        cols = [c for c in HOTEL_COLUMNS if c.key in ("option", "name", "rating", "price_per_night", "total_price")]
        t = doc.add_table(rows=1, cols=len(cols) + 2)
        t.style = "Light Grid Accent 1"
        for c, h in zip(t.rows[0].cells, [c.header for c in cols] + ["Amenities", "Link"]):
            _cell_text(c, h, bold=True)
        for row in report["hotels"]:
            cells = t.add_row().cells
            for cell, col in zip(cells, cols):
                _cell_text(cell, display_value(row, col))
            _cell_text(cells[-2], ", ".join(row.get("amenities_list") or []) or "n/a")
            url = safe_http_url(row.get("booking_url"))
            cells[-1].text = ""
            if url:
                _add_hyperlink(cells[-1].paragraphs[0], url, "Open")
            else:
                _cell_text(cells[-1], "n/a")

    if report.get("plan"):
        doc.add_heading("Suggested day-by-day plan", 1)
        doc.add_paragraph("Built from the returned attractions and weather. Places are grouped by proximity; opening hours are not checked.").runs[0].font.size = Pt(8)
        for day in report["plan"]:
            w = day.get("weather")
            wtxt = ""
            if w:
                wtxt = f" - {w.get('summary')}, {'n/a' if w.get('tmin') is None else round(w['tmin'])} to {'n/a' if w.get('tmax') is None else round(w['tmax'])} °C" + (" (typical)" if w.get("mode") == "typical" else "")
            p = doc.add_paragraph()
            p.add_run(f"{day['label']} - {day['date']}").bold = True
            p.add_run(clean_text(wtxt))
            for a in day["attractions"]:
                doc.add_paragraph(clean_text(a["name"] + (f" ({a['rating']:.1f})" if a.get("rating") is not None else "")), style="List Bullet")

    doc.add_heading(f"Tourist attractions ({len(report.get('attractions', []))})", 1)
    if not report.get("attractions"):
        doc.add_paragraph("No attractions were returned for this destination.")
    else:
        cols = [c for c in ATTRACTION_COLUMNS if c.key in ("option", "name", "type", "rating", "reviews")]
        t = doc.add_table(rows=1, cols=len(cols) + 1)
        t.style = "Light Grid Accent 1"
        for c, h in zip(t.rows[0].cells, [c.header for c in cols] + ["Map"]):
            _cell_text(c, h, bold=True)
        for row in report["attractions"]:
            cells = t.add_row().cells
            for cell, col in zip(cells, cols):
                _cell_text(cell, display_value(row, col))
            url = safe_http_url(row.get("map_url"))
            cells[-1].text = ""
            if url:
                _add_hyperlink(cells[-1].paragraphs[0], url, "Open")
            else:
                _cell_text(cells[-1], "n/a")

    doc.add_heading("Weather", 1)
    days = (report.get("weather") or {}).get("days") or []
    if not days:
        doc.add_paragraph("No weather data was available for these dates.")
    else:
        t = doc.add_table(rows=1, cols=5)
        t.style = "Light Grid Accent 1"
        for c, h in zip(t.rows[0].cells, ["Date", "Basis", "Conditions", "High / Low (°C)", "Rain"]):
            _cell_text(c, h, bold=True)
        for d_ in days:
            hi = "n/a" if d_.get("tmax") is None else f"{d_['tmax']:.0f}"
            lo = "n/a" if d_.get("tmin") is None else f"{d_['tmin']:.0f}"
            rain = f"{d_['rain_chance']}% chance" if d_.get("rain_chance") is not None else ("n/a" if d_.get("precip") is None else f"{d_['precip']:.1f} mm")
            cells = t.add_row().cells
            for cell, text in zip(cells, [d_["date"], d_.get("mode_label", ""), d_.get("summary", ""), f"{hi} / {lo}", rain]):
                _cell_text(cell, text)

    doc.add_heading("Data sources", 1)
    for s in report["sources"]:
        retrieved = f" (retrieved {s['retrieved']})" if s.get("retrieved") else ""
        doc.add_paragraph(clean_text(f"{s['type']}: {s['name']}{retrieved}"), style="List Bullet")
    doc.add_heading("Important limitations and notes", 1)
    for note in report["notes"]:
        doc.add_paragraph(clean_text(note), style="List Bullet")

    _add_page_number_footer(doc)
    doc.core_properties.title = clean_text(f"{report['title']} {req.get('origin')} to {req.get('destination')}")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
