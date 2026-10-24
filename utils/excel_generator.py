"""Excel (.xlsx) export built with openpyxl."""

from __future__ import annotations

import io
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .report_data import (
    ATTRACTION_COLUMNS,
    WEATHER_COLUMNS,
    FLIGHT_COLUMNS,
    HOTEL_COLUMNS,
    Column,
    clean_text,
    currency_symbol,
    safe_http_url,
)

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LINK_FONT = Font(color="0563C1", underline="single")
MAX_EXCEL_URL = 2000  # Excel refuses hyperlinks longer than ~2079 characters


def _money_format(currency: str) -> str:
    symbol = currency_symbol(currency)
    return f'"{symbol}"#,##0.00' if symbol else "#,##0.00"


def _set(cell, value: Any) -> None:
    """Write a value; text is forced to stay text so '=...' can never run as a formula."""
    if isinstance(value, str):
        cell.value = clean_text(value)
        cell.data_type = "s"
    else:
        cell.value = value


def _write_table(ws: Worksheet, columns: List[Column], rows: List[Dict[str, Any]], currency: str, empty_message: str) -> None:
    for c, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c, value=col.header)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = col.width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"

    if not rows:
        ws.cell(row=2, column=1, value=empty_message)
        return

    for r, row in enumerate(rows, start=2):
        for c, col in enumerate(columns, start=1):
            value = row.get(col.key)
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=col.kind == "text" and col.width >= 30)
            if value in (None, "", []):
                continue  # missing stays blank
            if col.kind == "url":
                url = safe_http_url(value)
                if url and len(url) <= MAX_EXCEL_URL:
                    cell.value = "Open link"
                    cell.hyperlink = url
                    cell.font = LINK_FONT
                elif url:
                    _set(cell, url)  # too long for a hyperlink: keep as plain text
                continue
            _set(cell, value)
            if col.kind == "money":
                cell.number_format = _money_format(currency)
            elif col.kind == "float":
                cell.number_format = "0.0"
            elif col.kind == "int":
                cell.number_format = "0"

    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"


def _write_pairs(ws: Worksheet, start_row: int, pairs: List[tuple]) -> int:
    row = start_row
    for label, value in pairs:
        a = ws.cell(row=row, column=1, value=label)
        a.font = Font(bold=True)
        a.alignment = Alignment(vertical="top")
        b = ws.cell(row=row, column=2)
        _set(b, "n/a" if value in (None, "", []) else value)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    return row


def build_excel(report: Dict[str, Any]) -> bytes:
    req = report["request"]
    currency = req.get("currency") or "USD"
    wb = Workbook()

    # --- Search Summary
    ws = wb.active
    ws.title = "Search Summary"
    ws["A1"] = report["title"]
    ws["A1"].font = TITLE_FONT
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 70
    travellers = f"{req.get('adults')} adult(s), {req.get('children')} child(ren)"
    if req.get("children_ages"):
        travellers += f" (ages {', '.join(str(a) for a in req['children_ages'])})"
    _write_pairs(
        ws,
        3,
        [
            ("Generated", report["generated_at_display"]),
            ("Origin", f"{req.get('origin')} ({req.get('origin_city')})" if req.get("origin_city") else req.get("origin")),
            ("Destination", f"{req.get('destination')} ({req.get('destination_city')})" if req.get("destination_city") else req.get("destination")),
            ("Hotel search city", req.get("hotel_search_city")),
            ("Departure date", req.get("departure_date")),
            ("Return date", req.get("return_date")),
            ("Nights", req.get("nights")),
            ("Travellers", travellers),
            ("Rooms", req.get("rooms")),
            ("Flight stops filter", req.get("flight_stops_filter")),
            ("Currency", currency),
            ("Flight options found", len(report["flights"])),
            ("Hotel options found", len(report["hotels"])),
        ],
    )

    # --- Flight Options
    _write_table(wb.create_sheet("Flight Options"), FLIGHT_COLUMNS, report["flights"], currency, "No flight options were returned for this search.")

    # --- Hotel Options
    _write_table(wb.create_sheet("Hotel Options"), HOTEL_COLUMNS, report["hotels"], currency, "No hotel options were returned for this search.")

    # --- Attractions / Weather / Trip Plan
    _write_table(wb.create_sheet("Attractions"), ATTRACTION_COLUMNS, report.get("attractions", []), currency, "No attractions were returned for this destination.")
    _write_table(wb.create_sheet("Weather"), WEATHER_COLUMNS, (report.get("weather") or {}).get("days", []), currency, "No weather data was available for these dates.")
    ws = wb.create_sheet("Trip Plan")
    for c, (h, w) in enumerate([("Day", 14), ("Date", 12), ("Weather", 36), ("Suggested attractions", 70)], start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    if not report.get("plan"):
        ws.cell(row=2, column=1, value="No plan could be built.")
    for r, day in enumerate(report.get("plan", []), start=2):
        w = day.get("weather") or {}
        wtxt = ""
        if w:
            wtxt = f"{w.get('summary')}, {'n/a' if w.get('tmin') is None else round(w['tmin'])} to {'n/a' if w.get('tmax') is None else round(w['tmax'])} °C" + (" (typical)" if w.get("mode") == "typical" else "")
        _set(ws.cell(row=r, column=1), day["label"])
        _set(ws.cell(row=r, column=2), day["date"])
        _set(ws.cell(row=r, column=3), wtxt)
        cell = ws.cell(row=r, column=4)
        _set(cell, "; ".join(a["name"] for a in day["attractions"]) or "Free time / travel")
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    # --- Recommendations (+ cost estimate)
    ws = wb.create_sheet("Recommendations")
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 18
    recs = report["recommendations"]
    ws["A1"] = "Recommendations"
    ws["A1"].font = TITLE_FONT
    pairs = [("Generated by", recs.get("source_label"))]
    for label, key, rows in (("Best flight", "best_flight", report["flights"]), ("Best hotel", "best_hotel", report["hotels"])):
        pick = recs.get(key)
        if pick:
            item = next((r for r in rows if r["option"] == pick["option"]), None)
            name = (item or {}).get("airline") or (item or {}).get("name") or ""
            pairs.append((label, f"Option {pick['option']}: {name}".strip()))
            pairs.append((f"{label}: reason", pick.get("reason")))
    pairs.append(("Summary", recs.get("summary")))
    for warning in recs.get("warnings") or []:
        pairs.append(("Note", warning))
    row = _write_pairs(ws, 3, pairs) + 1

    ws.cell(row=row, column=1, value="Estimated trip cost").font = TITLE_FONT
    row += 1
    scenarios = report["cost_estimate"]["scenarios"]
    if not scenarios:
        ws.cell(row=row, column=1, value="No estimate possible: no priced flight or hotel was returned.")
        row += 1
    for sc in scenarios:
        ws.cell(row=row, column=1, value=sc["scenario"]).font = Font(bold=True)
        row += 1
        for item in sc["items"]:
            _set(ws.cell(row=row, column=1), item["label"])
            _set(ws.cell(row=row, column=2), item["basis"])
            amount = ws.cell(row=row, column=3, value=item["amount"])
            amount.number_format = _money_format(currency)
            row += 1
        total = ws.cell(row=row, column=1, value="Total of items above" + ("" if sc["complete"] else " (INCOMPLETE: flight or hotel price missing)"))
        total.font = Font(bold=True)
        t = ws.cell(row=row, column=3, value=sc["total"])
        t.number_format = _money_format(currency)
        t.font = Font(bold=True)
        row += 2
    ws.cell(row=row, column=1, value="Not included in the estimate:").font = Font(bold=True)
    row += 1
    for text in report["cost_estimate"]["excluded"]:
        ws.cell(row=row, column=1, value="- " + text)
        row += 1

    # --- Trip Budget (only when the user entered a budget)
    budget = report.get("budget")
    if budget and budget.get("items"):
        ws = wb.create_sheet("Trip Budget")
        for c, (h, w) in enumerate([("Category", 18), ("Item", 44), ("Source", 52), (f"Amount ({currency})", 16)], start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
            ws.column_dimensions[get_column_letter(c)].width = w
        r = 2
        for item in budget["items"]:
            _set(ws.cell(row=r, column=1), item["category"]); _set(ws.cell(row=r, column=2), item["label"]); _set(ws.cell(row=r, column=3), item["basis"])
            amt = ws.cell(row=r, column=4, value=item["amount"]); amt.number_format = _money_format(currency)
            r += 1
        ws.cell(row=r, column=1, value="Total" + ("" if budget.get("complete") else " (INCOMPLETE: flight or hotel price missing)")).font = Font(bold=True)
        t = ws.cell(row=r, column=4, value=budget["total"]); t.number_format = _money_format(currency); t.font = Font(bold=True)
        r += 1
        conv = budget.get("converted")
        if conv:
            ws.cell(row=r, column=1, value=f"Total in {conv['currency']} (indicative rate 1 {currency} = {conv['rate']:.4f} {conv['currency']}, updated {conv.get('updated') or 'n/a'})")
            c = ws.cell(row=r, column=4, value=conv["total"]); c.number_format = _money_format(conv["currency"])

    # --- Sources and Notes
    ws = wb.create_sheet("Sources and Notes")
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 30
    for c, h in enumerate(["Type", "Source / provider", "Retrieved", "Link"], start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    r = 2
    for s in report["sources"]:
        _set(ws.cell(row=r, column=1), s["type"])
        _set(ws.cell(row=r, column=2), s["name"])
        if s.get("retrieved"):
            _set(ws.cell(row=r, column=3), s["retrieved"])
        url = safe_http_url(s.get("link"))
        if url and len(url) <= MAX_EXCEL_URL:
            link = ws.cell(row=r, column=4, value="Open link")
            link.hyperlink, link.font = url, LINK_FONT
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Important notes and limitations").font = TITLE_FONT
    r += 1
    for note in report["notes"]:
        cell = ws.cell(row=r, column=1)
        _set(cell, note)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 15 * max(1, len(note) // 110 + 1)
        r += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
